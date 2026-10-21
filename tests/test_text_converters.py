from __future__ import annotations

import pytest
from docx import Document
from odf.opendocument import OpenDocumentText
from odf.text import P
from pypdf import PdfReader

from file_converter.converters.text_converters import (
    docx_to_txt,
    odt_to_txt,
    txt_to_docx,
    txt_to_pdf,
)
from file_converter.exceptions import ConversionError


def test_txt_to_docx_preserves_lines_and_unicode(tmp_path):
    source = tmp_path / "input.txt"
    source.write_text("Zażółć gęślą jaźń\n\nŁódź, Польша\n", encoding="utf-8")
    target = tmp_path / "output.docx"

    txt_to_docx(source, target)
    document = Document(str(target))
    texts = [paragraph.text for paragraph in document.paragraphs]

    assert texts == ["Zażółć gęślą jaźń", "", "Łódź, Польша"]


def test_docx_to_txt_extracts_paragraphs_and_table(tmp_path):
    source = tmp_path / "input.docx"
    document = Document()
    document.add_paragraph("Zażółć gęślą jaźń")
    document.add_paragraph("Łódź")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Колонка 1"
    table.cell(0, 1).text = "Колонка 2"
    document.save(str(source))
    target = tmp_path / "output.txt"

    docx_to_txt(source, target)
    text = target.read_text(encoding="utf-8")

    assert "Zażółć gęślą jaźń" in text
    assert "Łódź" in text
    assert "Колонка 1\tКолонка 2" in text


def test_txt_to_pdf_contains_unicode_text(tmp_path):
    source = tmp_path / "input.txt"
    source.write_text("Zażółć gęślą jaźń\nŁódź\n" + "\n".join(f"linia {i}" for i in range(120)), encoding="utf-8")
    target = tmp_path / "output.pdf"

    txt_to_pdf(source, target)

    assert target.read_bytes().startswith(b"%PDF")
    reader = PdfReader(str(target))
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    assert "Zażółć" in text
    assert "Łódź" in text
    assert len(reader.pages) >= 2


def test_odt_to_txt_extracts_paragraphs(tmp_path):
    source = tmp_path / "input.odt"
    document = OpenDocumentText()
    document.text.addElement(P(text="Zażółć gęślą jaźń"))
    document.text.addElement(P(text="Łódź"))
    document.save(str(source))
    target = tmp_path / "output.txt"

    odt_to_txt(source, target)

    assert target.read_text(encoding="utf-8") == "Zażółć gęślą jaźń\nŁódź"


def test_invalid_docx_raises_conversion_error(tmp_path):
    source = tmp_path / "broken.docx"
    source.write_bytes(b"not a docx file")
    with pytest.raises(ConversionError, match="DOCX"):
        docx_to_txt(source, tmp_path / "out.txt")


def test_invalid_odt_raises_conversion_error(tmp_path):
    source = tmp_path / "broken.odt"
    source.write_bytes(b"not an odt file")
    with pytest.raises(ConversionError, match="ODT"):
        odt_to_txt(source, tmp_path / "out.txt")
