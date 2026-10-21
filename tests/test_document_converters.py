from __future__ import annotations

import subprocess
from pathlib import Path

import pytest
from docx import Document
from odf import teletype
from odf.opendocument import OpenDocumentText, load
from odf.text import P
from pypdf import PdfReader

from file_converter.converters import libreoffice
from file_converter.converters.document_converters import (
    docx_to_odt,
    docx_to_pdf,
    odt_to_docx,
    odt_to_pdf,
)
from file_converter.exceptions import ConversionError, DependencyMissingError, ExternalProcessError


@pytest.fixture
def docx_source(tmp_path):
    source = tmp_path / "dokument zażółć.docx"
    document = Document()
    document.add_paragraph("Zażółć gęślą jaźń")
    document.add_paragraph("Łódź — Польша")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Kolumna 1"
    table.cell(0, 1).text = "Kolumna 2"
    document.save(str(source))
    return source


@pytest.fixture
def odt_source(tmp_path):
    source = tmp_path / "dokument źródłowy.odt"
    document = OpenDocumentText()
    document.text.addElement(P(text="Zażółć gęślą jaźń"))
    document.text.addElement(P(text="Łódź — Польша"))
    document.save(str(source))
    return source


def _pdf_text(path: Path) -> str:
    return "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)


def _odt_text(path: Path) -> str:
    return teletype.extractText(load(str(path)).text)


def _docx_text(path: Path) -> str:
    return "\n".join(paragraph.text for paragraph in Document(str(path)).paragraphs)


def test_missing_libreoffice_raises(docx_source, tmp_path, monkeypatch):
    monkeypatch.setattr(libreoffice, "find_libreoffice", lambda: None)
    with pytest.raises(DependencyMissingError, match="LibreOffice"):
        docx_to_pdf(docx_source, tmp_path / "out.pdf")


def test_libreoffice_timeout(docx_source, tmp_path, monkeypatch):
    monkeypatch.setattr(libreoffice, "find_libreoffice", lambda: Path("/bin/echo"))

    def raise_timeout(*args, **kwargs):
        raise subprocess.TimeoutExpired(cmd=args[0], timeout=120)

    monkeypatch.setattr(libreoffice.subprocess, "run", raise_timeout)
    with pytest.raises(ExternalProcessError, match="не завершил работу"):
        docx_to_pdf(docx_source, tmp_path / "out.pdf")


def test_libreoffice_nonzero_exit(docx_source, tmp_path, monkeypatch):
    monkeypatch.setattr(libreoffice, "find_libreoffice", lambda: Path("/bin/echo"))
    completed = subprocess.CompletedProcess(["soffice"], 3, "", "conversion failed")
    monkeypatch.setattr(libreoffice.subprocess, "run", lambda *a, **k: completed)
    with pytest.raises(ExternalProcessError, match="кодом 3") as error:
        docx_to_pdf(docx_source, tmp_path / "out.pdf")
    assert "conversion failed" in error.value.stderr


def test_libreoffice_success_without_output_is_error(docx_source, tmp_path, monkeypatch):
    monkeypatch.setattr(libreoffice, "find_libreoffice", lambda: Path("/bin/echo"))
    completed = subprocess.CompletedProcess(["soffice"], 0, "ok", "")
    monkeypatch.setattr(libreoffice.subprocess, "run", lambda *a, **k: completed)
    with pytest.raises(ExternalProcessError, match="не создал"):
        docx_to_pdf(docx_source, tmp_path / "out.pdf")


REQUIRES_LIBREOFFICE = pytest.mark.skipif(
    not libreoffice.is_libreoffice_available(),
    reason="LibreOffice не установлен в этой среде",
)


@REQUIRES_LIBREOFFICE
def test_real_docx_to_pdf_preserves_text(docx_source, tmp_path):
    original = docx_source.read_bytes()
    target = tmp_path / "wynik zażółć.pdf"

    docx_to_pdf(docx_source, target)

    assert target.read_bytes().startswith(b"%PDF")
    text = _pdf_text(target)
    assert "Zażółć" in text
    assert "Łódź" in text
    assert docx_source.read_bytes() == original


@REQUIRES_LIBREOFFICE
def test_real_odt_to_pdf_preserves_text(odt_source, tmp_path):
    original = odt_source.read_bytes()
    target = tmp_path / "wynik źródłowy.pdf"

    odt_to_pdf(odt_source, target)

    assert target.read_bytes().startswith(b"%PDF")
    text = _pdf_text(target)
    assert "Zażółć" in text
    assert "Łódź" in text
    assert odt_source.read_bytes() == original


@REQUIRES_LIBREOFFICE
def test_real_docx_to_odt_preserves_text(docx_source, tmp_path):
    original = docx_source.read_bytes()
    target = tmp_path / "wynik zażółć.odt"

    docx_to_odt(docx_source, target)

    text = _odt_text(target)
    assert "Zażółć" in text
    assert "Łódź" in text
    assert docx_source.read_bytes() == original


@REQUIRES_LIBREOFFICE
def test_real_odt_to_docx_preserves_text(odt_source, tmp_path):
    original = odt_source.read_bytes()
    target = tmp_path / "wynik źródłowy.docx"

    odt_to_docx(odt_source, target)

    text = _docx_text(target)
    assert "Zażółć" in text
    assert "Łódź" in text
    assert odt_source.read_bytes() == original
