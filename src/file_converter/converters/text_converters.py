from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from odf import teletype
from odf.opendocument import load as load_odt
from odf.table import TableRow
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

from ..config import DEFAULT_ENCODING
from ..exceptions import ConversionError, DependencyMissingError

_FONT_CANDIDATES = (
    "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
    "/Library/Fonts/Arial Unicode.ttf",
    "/System/Library/Fonts/Supplemental/Arial.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/TTF/DejaVuSans.ttf",
    "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
    r"C:\Windows\Fonts\arial.ttf",
    r"C:\Windows\Fonts\DejaVuSans.ttf",
)

_FONT_NAME = "UFCUnicode"
_font_registered = False


def read_text_file(source: Path) -> str:
    source = Path(source)
    try:
        return source.read_text(encoding="utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ConversionError(
            "The text file is not valid UTF-8.",
            "Save the file as UTF-8 and try again.",
        ) from exc
    except OSError as exc:
        raise ConversionError(f"Could not read the file: {source}", str(exc)) from exc


def _split_lines(text: str) -> list[str]:
    return text.replace("\r\n", "\n").replace("\r", "\n").split("\n")


def txt_to_docx(source: Path, target: Path) -> Path:
    text = read_text_file(source)
    document = Document()
    lines = _split_lines(text)
    if lines and lines[-1] == "":
        lines = lines[:-1]
    if not lines:
        document.add_paragraph("")
    for line in lines:
        document.add_paragraph(line)
    try:
        document.save(str(target))
    except OSError as exc:
        raise ConversionError(f"Could not write the DOCX file: {target}", str(exc)) from exc
    return Path(target)


def _iter_blocks(document: Document):
    body = document.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def docx_to_txt(source: Path, target: Path) -> Path:
    source = Path(source)
    try:
        document = Document(str(source))
    except Exception as exc:
        raise ConversionError(
            f"Could not open the DOCX file: {source}",
            "The file is damaged or is not a Word document.",
        ) from exc

    lines: list[str] = []
    try:
        for block in _iter_blocks(document):
            if isinstance(block, Paragraph):
                lines.append(block.text)
            else:
                for row in block.rows:
                    cells = [
                        " ".join(p.text for p in cell.paragraphs).strip()
                        for cell in row.cells
                    ]
                    lines.append("\t".join(cells))
    except Exception as exc:
        raise ConversionError(
            "Could not extract text from the DOCX file.",
            "The document contains an unreadable structure.",
        ) from exc

    text = "\n".join(lines)
    try:
        Path(target).write_text(text, encoding=DEFAULT_ENCODING)
    except OSError as exc:
        raise ConversionError(f"Could not write the TXT file: {target}", str(exc)) from exc
    return Path(target)


def _register_unicode_font() -> str:
    global _font_registered
    if _font_registered:
        return _FONT_NAME
    for candidate in _FONT_CANDIDATES:
        if Path(candidate).is_file():
            try:
                pdfmetrics.registerFont(TTFont(_FONT_NAME, candidate))
                _font_registered = True
                return _FONT_NAME
            except Exception as exc:
                raise DependencyMissingError(
                    f"Could not register the font {candidate}.",
                    str(exc),
                ) from exc
    raise DependencyMissingError(
        "No Unicode font was found for PDF creation.",
        "Install DejaVu Sans, Arial, or Liberation Sans and try again.",
    )


def _wrap_line(text: str, font: str, size: float, max_width: float) -> list[str]:
    if not text:
        return [""]
    words = text.split(" ")
    lines: list[str] = []
    current = ""
    for word in words:
        candidate = word if not current else f"{current} {word}"
        if pdfmetrics.stringWidth(candidate, font, size) <= max_width:
            current = candidate
            continue
        if current:
            lines.append(current)
        while pdfmetrics.stringWidth(word, font, size) > max_width:
            cut = len(word)
            while cut > 1 and pdfmetrics.stringWidth(word[:cut], font, size) > max_width:
                cut -= 1
            lines.append(word[:cut])
            word = word[cut:]
        current = word
    lines.append(current)
    return lines


def txt_to_pdf(source: Path, target: Path) -> Path:
    text = read_text_file(source)
    font = _register_unicode_font()
    size = 11
    leading = 15
    width, height = A4
    margin = 56
    usable_width = width - 2 * margin
    top = height - margin

    pdf = canvas.Canvas(str(target), pagesize=A4)
    pdf.setTitle(Path(source).stem)
    pdf.setFont(font, size)
    y = top
    for paragraph in _split_lines(text):
        for line in _wrap_line(paragraph, font, size, usable_width):
            if y < margin:
                pdf.showPage()
                pdf.setFont(font, size)
                y = top
            if line:
                pdf.drawString(margin, y, line)
            y -= leading
    pdf.save()
    return Path(target)


def _collect_odt_blocks(node, lines: list[str]) -> None:
    qname = getattr(node, "qname", (None, ""))
    local = qname[1] if qname else ""
    if local in {"p", "h"}:
        lines.append(teletype.extractText(node))
        return
    if local == "table":
        for row in node.getElementsByType(TableRow):
            cells = []
            for cell in row.childNodes:
                if getattr(cell, "qname", ("", ""))[1] != "table-cell":
                    continue
                cell_lines: list[str] = []
                for child in cell.childNodes:
                    _collect_odt_blocks(child, cell_lines)
                cells.append(" ".join(part for part in cell_lines if part).strip())
            lines.append("\t".join(cells))
        return
    for child in getattr(node, "childNodes", []):
        _collect_odt_blocks(child, lines)


def odt_to_txt(source: Path, target: Path) -> Path:
    source = Path(source)
    try:
        document = load_odt(str(source))
    except Exception as exc:
        raise ConversionError(
            f"Could not open the ODT file: {source}",
            "The file is damaged or is not an OpenDocument document.",
        ) from exc

    lines: list[str] = []
    try:
        _collect_odt_blocks(document.text, lines)
    except Exception as exc:
        raise ConversionError(
            "Could not extract text from the ODT file.",
            "The document contains an unreadable structure.",
        ) from exc

    try:
        Path(target).write_text("\n".join(lines), encoding=DEFAULT_ENCODING)
    except OSError as exc:
        raise ConversionError(f"Could not write the TXT file: {target}", str(exc)) from exc
    return Path(target)
